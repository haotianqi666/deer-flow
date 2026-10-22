import json
import subprocess
import sys
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

SCRIPT = (
    Path(__file__).resolve().parents[2]
    / "skills/public/ppt-generation/scripts/generate.py"
)


def _write_slide(path, kind):
    if kind in ("RGBA", "LA"):
        image = Image.new(kind, (160, 90), (0, 0, 0, 0) if kind == "RGBA" else (0, 0))
        image.paste(
            (80, 120, 160, 128) if kind == "RGBA" else (80, 128), (40, 0, 80, 90)
        )
        image.paste(
            (80, 120, 160, 255) if kind == "RGBA" else (80, 255), (80, 0, 160, 90)
        )
    elif kind.startswith("palette"):
        image = Image.new("P", (160, 90), 0)
        image.putpalette([0, 0, 0, 80, 120, 160, 160, 120, 80] + [0] * 759)
        image.paste(1, (40, 0, 80, 90))
        image.paste(2, (80, 0, 160, 90))
        image.info["transparency"] = (
            0 if kind == "palette-index" else bytes([0, 128, 255])
        )
    elif kind in ("RGB-key", "L-key"):
        mode = kind.split("-")[0]
        image = Image.new(mode, (160, 90), 0)
        image.paste((80, 120, 160) if mode == "RGB" else 80, (80, 0, 160, 90))
        image.info["transparency"] = (0, 0, 0) if mode == "RGB" else 0
    else:
        mode = "RGB" if kind == "JPEG" else kind
        image = Image.new(mode, (160, 90), (80, 120, 160) if mode == "RGB" else 80)
        if mode == "P":
            image.putpalette([value for value in range(256) for _ in range(3)])
    image.save(path)


def _compose(tmp_path, image_path, aspect_ratio="16:9"):
    plan = tmp_path / "plan.json"
    plan.write_text(
        json.dumps(
            {
                "aspect_ratio": aspect_ratio,
                "slides": [
                    {
                        "title": "A slide",
                        "subtitle": "A subtitle",
                        "key_points": ["A point"],
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    output = tmp_path / "deck.pptx"
    result = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--plan-file",
            str(plan),
            "--slide-images",
            str(image_path),
            "--output-file",
            str(output),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=30,
        check=False,
    )
    assert output.exists(), result.stdout + result.stderr
    assert result.returncode == 0, result.stderr
    assert "Successfully generated presentation with 1 slides" in result.stdout
    presentation = Presentation(output)
    assert len(presentation.slides) == 1
    assert presentation.slide_width == Inches(13.333 if aspect_ratio == "16:9" else 10)
    assert presentation.slide_height == Inches(7.5)
    slide = presentation.slides[0]
    assert "Title: A slide" in slide.notes_slide.notes_text_frame.text
    assert "Subtitle: A subtitle" in slide.notes_slide.notes_text_frame.text
    assert "A point" in slide.notes_slide.notes_text_frame.text
    return Image.open(BytesIO(slide.shapes[0].image.blob))


@pytest.mark.parametrize("aspect_ratio", ["16:9", "4:3"])
@pytest.mark.parametrize(
    "kind", ["RGBA", "LA", "palette-index", "palette-alpha", "RGB-key", "L-key"]
)
def test_ppt_cli_preserves_transparent_slide(tmp_path, kind, aspect_ratio):
    source = tmp_path / "slide.png"
    _write_slide(source, kind)
    original_bytes = source.read_bytes()

    with (
        Image.open(source) as expected,
        _compose(tmp_path, source, aspect_ratio) as actual,
    ):
        assert actual.size == expected.size
        # Compare the embedded payload, including partial alpha, not just the file type.
        assert actual.convert("RGBA").tobytes() == expected.convert("RGBA").tobytes()
    assert source.read_bytes() == original_bytes


@pytest.mark.parametrize("kind", ["RGB", "L", "P", "JPEG"])
def test_ppt_cli_keeps_opaque_image_behavior(tmp_path, kind):
    source = tmp_path / ("slide.jpg" if kind == "JPEG" else "slide.png")
    _write_slide(source, kind)
    original_bytes = source.read_bytes()

    with Image.open(source) as expected, _compose(tmp_path, source) as actual:
        assert actual.format == "JPEG"
        assert actual.size == expected.size
        expected_pixel = expected.convert("RGB").getpixel((80, 45))
        actual_pixel = actual.convert("RGB").getpixel((80, 45))
        assert all(
            abs(before - after) <= 2
            for before, after in zip(expected_pixel, actual_pixel)
        )
    assert source.read_bytes() == original_bytes
